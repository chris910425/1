"""Put a recorded motion clip on a full-bleed PowerPoint slide that starts
automatically and loops until the slide changes.

    python3 tools/record-motion/to_pptx.py clips/neon.mp4 [--out clips/neon.pptx] [--append deck.pptx]

Uses the clip's .png (written by record.js) as the poster frame. With --append the
slide is added to the end of an existing deck instead of a new 16:9 file.
Needs python-pptx.
"""
import argparse, copy, pathlib
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

P = "http://schemas.openxmlformats.org/presentationml/2006/main"

TIMING = """<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{SPID}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="80000" mute="1"><p:cTn id="7" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{SPID}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("clip")
    ap.add_argument("--out")
    ap.add_argument("--append")
    a = ap.parse_args()
    clip = pathlib.Path(a.clip)
    poster = clip.with_suffix(".png")
    prs = Presentation(a.append) if a.append else Presentation()
    if not a.append:
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9
    blank = min(prs.slide_layouts, key=lambda l: len(l.placeholders))
    slide = prs.slides.add_slide(blank)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    movie = slide.shapes.add_movie(str(clip), 0, 0, prs.slide_width, prs.slide_height,
                                   poster_frame_image=str(poster) if poster.exists() else None, mime_type="video/mp4")
    movie.name = f"motion: {clip.stem}"
    # replace python-pptx's click-to-play timing with autoplay + loop
    sld = slide._element
    for old in sld.findall(f"{{{P}}}timing"):
        sld.remove(old)
    timing = etree.fromstring(TIMING.replace("{P}", P).replace("{SPID}", str(movie.shape_id)))
    ext = sld.find(f"{{{P}}}extLst")
    (ext.addprevious(timing) if ext is not None else sld.append(timing))
    out = pathlib.Path(a.out or (a.append if a.append else clip.with_suffix(".pptx")))
    prs.save(out)
    print(f"saved {out} (slide {len(prs.slides)}, video autoplays and loops, muted)")


if __name__ == "__main__":
    main()
